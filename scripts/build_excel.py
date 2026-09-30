import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment
from data import rows

F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_sub = Font(name=F, size=10, italic=True, color="666666")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
fill_head = PatternFill("solid", fgColor="1F4E78")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_band = PatternFill("solid", fgColor="F2F2F2")
fill_key = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()

# ---------------------------------------------------------------- 说明
ws0 = wb.active
ws0.title = "说明"
ws0["A1"] = "国创光谷上城 二手房成交数据整理（卖房参考）"
ws0["A1"].font = f_title
notes = [
    ("数据来源", "用户提供的 20 张经纪 App 成交记录截图（截图时间 2026-09-30），共 80 条，成交日期 2025-02-16 至 2026-09-10。"),
    ("字段口径", "户型「3-2-1-2」= 3室2厅1厨2卫；楼层仅有高/中/低三段，截图未给具体层数；单价为平台显示值，另列按 总价÷面积 的核算值做校验（3 条有 2~7 元尾差，属平台四舍五入或面积小数位差异）。"),
    ("状态", "「已成交」= 已签约未过户（2 条，均为 2026 年 9 月前后）；其余为「已过户」。第 11 张截图顶部 1 条的标签被截断，按已过户记。"),
    ("可能遗漏", "截图之间疑似有未截到的记录：2025-07-20 至 2025-09-09 之间（2025 年 8 月无记录）、2025-04-06 至 2025-05-25 之间；2026 年 2 月无记录（可能是春节淡季）。统计结论按 80 条样本，存在样本偏差。"),
    ("首图观感", "按列表缩略图目测：有装修 / 毛坯 / 空房(不明) / 户型图 / 无图。仅供参考，不代表实际交付标准。"),
    ("颜色约定", "蓝字黄底 = 可修改的假设或参数；黑字 = 公式；绿字 = 引用其他工作表。所有统计与测算均为公式，修改明细或参数后自动重算。"),
    ("工作表", "成交明细 → 季度走势 → 分面积段对比 → 本案估值（110㎡/中楼层/南北通透）→ 持有vs出售测算"),
    ("免责声明", "本表仅为基于公开成交记录的整理与测算，不构成投资、税务或法律建议；税费、贷款、政策以武汉市不动产登记与税务部门最新口径为准。"),
]
r = 3
for k, v in notes:
    ws0.cell(row=r, column=1, value=k).font = f_bold
    c = ws0.cell(row=r, column=2, value=v)
    c.font = f_base
    c.alignment = left
    ws0.row_dimensions[r].height = 42
    r += 1
ws0.column_dimensions["A"].width = 14
ws0.column_dimensions["B"].width = 110
for rr in range(3, r):
    ws0.cell(row=rr, column=1).alignment = Alignment(vertical="top")

# ---------------------------------------------------------------- 成交明细
ws = wb.create_sheet("成交明细")
headers = ["序号", "截图编号", "状态", "成交日期", "年份", "季度", "半年", "户型(原始)", "室", "厅", "厨", "卫",
           "户型", "建筑面积(㎡)", "面积段", "楼层段", "总层数", "朝向", "南北通透", "首图观感(目测)",
           "成交总价(万元)", "平台单价(元/㎡)", "核算单价(元/㎡)", "单价尾差", "本案可比样本"]
for i, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = f_head
    c.fill = fill_head
    c.alignment = center
    c.border = border

N = len(rows)
for idx, (img, status, date, layout, area, floor, tf, orient, total, unit, decor) in enumerate(rows, 1):
    r = idx + 1
    rm, ht, kt, wt = [int(x) for x in layout.split("-")]
    vals = {
        1: idx, 2: img, 3: status, 4: dt.datetime.strptime(date, "%Y-%m-%d"),
        5: f"=YEAR(D{r})",
        6: f'=YEAR(D{r})&"Q"&ROUNDUP(MONTH(D{r})/3,0)',
        7: f'=YEAR(D{r})&"H"&IF(MONTH(D{r})<=6,1,2)',
        8: layout, 9: rm, 10: ht, 11: kt, 12: wt,
        13: f'=I{r}&"室"&J{r}&"厅"&L{r}&"卫"',
        14: area,
        15: f'=IF(N{r}<60,"60㎡以下",IF(N{r}<90,"75-77㎡",IF(N{r}<100,"94-97㎡",IF(N{r}<120,"106-117㎡",IF(N{r}<150,"124㎡","150㎡以上")))))',
        16: floor, 17: tf, 18: orient,
        19: f'=IF(AND(ISNUMBER(SEARCH("南",R{r})),ISNUMBER(SEARCH("北",R{r}))),"是","否")',
        20: decor, 21: total, 22: unit,
        23: f"=ROUND(U{r}*10000/N{r},0)",
        24: f"=W{r}-V{r}",
        25: f"=IF(AND(I{r}=3,N{r}>='本案估值'!$C$7,N{r}<='本案估值'!$C$8,D{r}>='本案估值'!$C$9),\"是\",\"否\")",
    }
    for col, v in vals.items():
        c = ws.cell(row=r, column=col, value=v)
        c.font = f_base
        c.border = border
        c.alignment = center
        if idx % 2 == 0:
            c.fill = fill_band
    ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=14).number_format = "0.00"
    ws.cell(row=r, column=21).number_format = "0.0"
    for col in (22, 23):
        ws.cell(row=r, column=col).number_format = "#,##0"
    ws.cell(row=r, column=24).number_format = "0;-0;-"

LAST = N + 1
widths = [6, 8, 16, 11, 7, 9, 8, 10, 5, 5, 5, 5, 10, 11, 10, 7, 7, 10, 8, 13, 12, 12, 12, 8, 11]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{LAST}"
ws.row_dimensions[1].height = 30
ws["T1"].comment = Comment("按列表缩略图目测，仅供参考。", "整理")
ws["Y1"].comment = Comment("条件在「本案估值」C7:C9 设置：3室、面积区间、起始日期。", "整理")

D = "成交明细"
rng = lambda col: f"'{D}'!${col}$2:${col}${LAST}"

# ---------------------------------------------------------------- 季度走势
wq = wb.create_sheet("季度走势")
wq["A1"] = "季度成交走势（全样本 & 106-117㎡ 三房）"
wq["A1"].font = f_title
wq["A2"] = "注：部分季度样本量很小（≤7 套），单套极端值会明显拉动均价，中位数更稳健；2025Q3 仅含 7 月与 9 月记录。"
wq["A2"].font = f_sub
qh = ["季度", "成交套数", "面积加权均价(元/㎡)", "单价中位数(元/㎡)", "最低单价", "最高单价",
      "较2025Q1 (加权均价)", "106-117㎡ 三房套数", "106-117㎡ 三房中位数", "106-117㎡ 较2025Q1"]
for i, h in enumerate(qh, 1):
    c = wq.cell(row=4, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
quarters = ["2025Q1", "2025Q2", "2025Q3", "2025Q4", "2026Q1", "2026Q2", "2026Q3"]
for j, q in enumerate(quarters):
    r = 5 + j
    wq.cell(row=r, column=1, value=q)
    wq.cell(row=r, column=2, value=f"=COUNTIFS({rng('F')},A{r})")
    wq.cell(row=r, column=3, value=f"=IFERROR(SUMIFS({rng('U')},{rng('F')},A{r})*10000/SUMIFS({rng('N')},{rng('F')},A{r}),\"-\")")
    wq.cell(row=r, column=4).value = ArrayFormula(f"D{r}", f"=IFERROR(MEDIAN(IF({rng('F')}=A{r},{rng('W')})),\"-\")")
    wq.cell(row=r, column=5, value=f"=IFERROR(_xlfn.MINIFS({rng('W')},{rng('F')},A{r}),\"-\")")
    wq.cell(row=r, column=6, value=f"=IFERROR(_xlfn.MAXIFS({rng('W')},{rng('F')},A{r}),\"-\")")
    wq.cell(row=r, column=7, value=f"=IFERROR(C{r}/$C$5-1,\"-\")")
    wq.cell(row=r, column=8, value=f"=COUNTIFS({rng('F')},A{r},{rng('I')},3,{rng('O')},\"106-117㎡\")")
    wq.cell(row=r, column=9).value = ArrayFormula(f"I{r}", f"=IFERROR(MEDIAN(IF(({rng('F')}=A{r})*({rng('I')}=3)*({rng('O')}=\"106-117㎡\"),{rng('W')})),\"-\")")
    wq.cell(row=r, column=10, value=f"=IFERROR(I{r}/$I$5-1,\"-\")")
    for col in range(1, 11):
        c = wq.cell(row=r, column=col)
        c.font = f_base; c.border = border; c.alignment = center
    for col in (3, 4, 5, 6, 9):
        wq.cell(row=r, column=col).number_format = "#,##0"
    for col in (7, 10):
        wq.cell(row=r, column=col).number_format = "0.0%;-0.0%;-"

# half-year block
hr0 = 14
wq.cell(row=hr0 - 1, column=1, value="半年度汇总").font = f_bold
hh = ["半年", "成交套数", "面积加权均价(元/㎡)", "单价中位数(元/㎡)", "较2025H1 (中位数)"]
for i, h in enumerate(hh, 1):
    c = wq.cell(row=hr0, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
halves = ["2025H1", "2025H2", "2026H1", "2026H2"]
for j, hname in enumerate(halves):
    r = hr0 + 1 + j
    wq.cell(row=r, column=1, value=hname)
    wq.cell(row=r, column=2, value=f"=COUNTIFS({rng('G')},A{r})")
    wq.cell(row=r, column=3, value=f"=IFERROR(SUMIFS({rng('U')},{rng('G')},A{r})*10000/SUMIFS({rng('N')},{rng('G')},A{r}),\"-\")")
    wq.cell(row=r, column=4).value = ArrayFormula(f"D{r}", f"=IFERROR(MEDIAN(IF({rng('G')}=A{r},{rng('W')})),\"-\")")
    wq.cell(row=r, column=5, value=f"=IFERROR(D{r}/$D${hr0+1}-1,\"-\")")
    for col in range(1, 6):
        c = wq.cell(row=r, column=col)
        c.font = f_base; c.border = border; c.alignment = center
    for col in (3, 4):
        wq.cell(row=r, column=col).number_format = "#,##0"
    wq.cell(row=r, column=5).number_format = "0.0%;-0.0%;-"
wq.cell(row=hr0 + 5, column=1, value="2026H2 目前只含 7-9 月。").font = f_sub

for i, w in enumerate([10, 10, 16, 16, 11, 11, 14, 14, 16, 14], 1):
    wq.column_dimensions[get_column_letter(i)].width = w
wq.row_dimensions[4].height = 32

ch = LineChart()
ch.title = "季度成交单价走势（元/㎡）"
ch.y_axis.title = "元/㎡"
ch.x_axis.title = "季度"
ch.height = 8.5
ch.width = 18
data = Reference(wq, min_col=3, min_row=4, max_row=11)
ch.add_data(data, titles_from_data=True)
data2 = Reference(wq, min_col=9, min_row=4, max_row=11)
ch.add_data(data2, titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=5, max_row=11))
ch.y_axis.scaling.min = 9000
ch.y_axis.scaling.max = 15000
ch.y_axis.delete = False
ch.x_axis.delete = False
for s_ in ch.series:
    s_.smooth = False
wq.add_chart(ch, "A22")

bc = BarChart()
bc.title = "季度成交套数"
bc.height = 8.5
bc.width = 12
bc.add_data(Reference(wq, min_col=2, min_row=4, max_row=11), titles_from_data=True)
bc.set_categories(Reference(wq, min_col=1, min_row=5, max_row=11))
bc.y_axis.delete = False
bc.x_axis.delete = False
bc.legend = None
wq.add_chart(bc, "H22")

# ---------------------------------------------------------------- 分面积段对比
wa = wb.create_sheet("分面积段对比")
wa["A1"] = "分面积段成交单价对比（中位数，元/㎡）"
wa["A1"].font = f_title
wa["A2"] = "套数与中位数分列；样本少于 3 套的格子参考价值有限。"
wa["A2"].font = f_sub
bins = ["60㎡以下", "75-77㎡", "94-97㎡", "106-117㎡", "124㎡", "150㎡以上"]
ah = ["面积段"]
for hname in halves:
    ah += [f"{hname} 套数", f"{hname} 中位数"]
ah += ["2026H2 较 2025H1"]
for i, h in enumerate(ah, 1):
    c = wa.cell(row=4, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
for j, b in enumerate(bins):
    r = 5 + j
    wa.cell(row=r, column=1, value=b)
    col = 2
    for hname in halves:
        wa.cell(row=r, column=col, value=f"=COUNTIFS({rng('O')},$A{r},{rng('G')},\"{hname}\")")
        cl = get_column_letter(col + 1)
        wa.cell(row=r, column=col + 1).value = ArrayFormula(
            f"{cl}{r}", f"=IFERROR(MEDIAN(IF(({rng('O')}=$A{r})*({rng('G')}=\"{hname}\"),{rng('W')})),\"-\")")
        wa.cell(row=r, column=col + 1).number_format = "#,##0"
        col += 2
    wa.cell(row=r, column=col, value=f"=IFERROR(I{r}/C{r}-1,\"-\")").number_format = "0.0%;-0.0%;-"
    for cc in range(1, col + 1):
        c = wa.cell(row=r, column=cc)
        c.font = f_base; c.border = border; c.alignment = center
for i in range(1, len(ah) + 1):
    wa.column_dimensions[get_column_letter(i)].width = 12
wa.row_dimensions[4].height = 32

wa["A13"] = "朝向与楼层（仅 106-117㎡ 三房，全时段）"
wa["A13"].font = f_bold
oh = ["维度", "取值", "套数", "单价中位数(元/㎡)"]
for i, h in enumerate(oh, 1):
    c = wa.cell(row=14, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
dims = [("南北通透", "是", "S"), ("南北通透", "否", "S"), ("楼层段", "高", "P"), ("楼层段", "中", "P"), ("楼层段", "低", "P")]
for j, (dname, val, colk) in enumerate(dims):
    r = 15 + j
    wa.cell(row=r, column=1, value=dname)
    wa.cell(row=r, column=2, value=val)
    wa.cell(row=r, column=3, value=f"=COUNTIFS({rng('O')},\"106-117㎡\",{rng('I')},3,{rng(colk)},B{r})")
    wa.cell(row=r, column=4).value = ArrayFormula(
        f"D{r}", f"=IFERROR(MEDIAN(IF(({rng('O')}=\"106-117㎡\")*({rng('I')}=3)*({rng(colk)}=B{r}),{rng('W')})),\"-\")")
    wa.cell(row=r, column=4).number_format = "#,##0"
    for cc in range(1, 5):
        c = wa.cell(row=r, column=cc)
        c.font = f_base; c.border = border; c.alignment = center
wa["A21"] = "注：全时段混合了 2025 年高价期与 2026 年低价期，朝向/楼层差异可能被成交时间差掩盖，只作方向参考。"
wa["A21"].font = f_sub

# ---------------------------------------------------------------- 本案估值
wv = wb.create_sheet("本案估值")
wv["A1"] = "本案估值：110㎡ · 中间楼栋 14/33 层（中楼层）· 南北通透"
wv["A1"].font = f_title
wv["A2"] = "黄底蓝字为可调参数。可比样本 = 3室 + 面积区间 + 起始日期之后成交（在「成交明细」Y 列标记）。"
wv["A2"].font = f_sub

def inp(r, label, val, fmt=None, note=None):
    wv.cell(row=r, column=2, value=label).font = f_base
    c = wv.cell(row=r, column=3, value=val)
    c.font = f_input; c.fill = fill_input; c.border = border; c.alignment = center
    if fmt: c.number_format = fmt
    if note:
        n = wv.cell(row=r, column=4, value=note); n.font = f_sub; n.alignment = left

def out(r, label, formula, fmt=None, note=None, key=False):
    wv.cell(row=r, column=2, value=label).font = f_bold if key else f_base
    c = wv.cell(row=r, column=3, value=formula)
    c.font = f_bold if key else f_base; c.border = border; c.alignment = center
    if key: c.fill = fill_key
    if fmt: c.number_format = fmt
    if note:
        n = wv.cell(row=r, column=4, value=note); n.font = f_sub; n.alignment = left

wv.cell(row=4, column=2, value="一、本案与可比条件").font = f_bold
inp(5, "本案建筑面积(㎡)", 110, "0.00", "用户给定")
inp(6, "本案楼层", "中（14/33）", None, "用户给定；小区中间楼栋，通常噪音与视野较好")
inp(7, "可比样本：最小面积(㎡)", 105, "0", "取 106-117㎡ 大三房（小区 109.99/113.37/115.97/116.04 等户型）")
inp(8, "可比样本：最大面积(㎡)", 117, "0")
inp(9, "可比样本：起始日期", dt.datetime(2026, 3, 1), "yyyy-mm-dd", "只取近约 6 个月，反映当前价格水平")
inp(10, "近期窗口起始日期", dt.datetime(2026, 7, 1), "yyyy-mm-dd", "用于单独观察最近一个季度")

wv.cell(row=12, column=2, value="二、可比样本统计").font = f_bold
Y = rng("Y"); W = rng("W"); Dd = rng("D")
out(13, "可比样本套数", f"=COUNTIFS({Y},\"是\")", "0")
out(14, "可比单价均值(元/㎡)", f"=IFERROR(AVERAGEIFS({W},{Y},\"是\"),\"-\")", "#,##0")
wv.cell(row=15, column=2, value="可比单价中位数(元/㎡)").font = f_base
wv["C15"] = ArrayFormula("C15", f"=MEDIAN(IF({Y}=\"是\",{W}))")
wv.cell(row=16, column=2, value="可比单价 25 分位(元/㎡)").font = f_base
wv["C16"] = ArrayFormula("C16", f"=PERCENTILE(IF({Y}=\"是\",{W}),0.25)")
wv.cell(row=17, column=2, value="可比单价 75 分位(元/㎡)").font = f_base
wv["C17"] = ArrayFormula("C17", f"=PERCENTILE(IF({Y}=\"是\",{W}),0.75)")
out(18, "近期窗口可比套数", f"=COUNTIFS({Y},\"是\",{Dd},\">=\"&C10)", "0")
out(19, "近期窗口可比均价(元/㎡)", f"=IFERROR(AVERAGEIFS({W},{Y},\"是\",{Dd},\">=\"&C10),\"-\")", "#,##0")
for r in (15, 16, 17):
    c = wv.cell(row=r, column=3); c.font = f_base; c.border = border; c.alignment = center; c.number_format = "#,##0"

wv.cell(row=21, column=2, value="三、个性化调整（假设，可改）").font = f_bold
inp(22, "南北通透调整", 0.01, "0.0%", "假设值：样本里南北通透与纯南向差异不显著，给小幅溢价")
inp(23, "中楼层/中间楼栋调整", 0.01, "0.0%", "假设值：可比样本含较多高/低楼层")
inp(24, "装修状况调整", 0.0, "0.0%", "请按实际填写：精装保养好可 +2%~+4%，毛坯/老旧 -3%~-5%")
out(25, "合计调整", "=C22+C23+C24", "0.0%")

wv.cell(row=27, column=2, value="四、估值结果").font = f_bold
hdr = ["", "情形", "单价(元/㎡)", "总价(万元)"]
for i, h in enumerate(hdr[1:], 2):
    c = wv.cell(row=28, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
scen = [("保守（25 分位）", "C16"), ("中性（中位数）", "C15"), ("乐观（75 分位）", "C17"), ("参考：近期窗口均价", "C19")]
for j, (lab, ref) in enumerate(scen):
    r = 29 + j
    wv.cell(row=r, column=2, value=lab)
    wv.cell(row=r, column=3, value=f"=IFERROR(ROUND({ref}*(1+$C$25),0),\"-\")").number_format = "#,##0"
    wv.cell(row=r, column=4, value=f"=IFERROR(C{r}*$C$5/10000,\"-\")").number_format = "0.0"
    for cc in (2, 3, 4):
        c = wv.cell(row=r, column=cc); c.font = f_base; c.border = border; c.alignment = center
    if j == 1:
        for cc in (2, 3, 4):
            wv.cell(row=r, column=cc).fill = fill_key
            wv.cell(row=r, column=cc).font = f_bold

wv.cell(row=34, column=2, value="五、挂牌策略参考").font = f_bold
inp(35, "挂牌议价空间", 0.04, "0.0%", "假设值：挂牌价高于目标成交价的幅度，留给买方砍价")
out(36, "建议挂牌价(万元)", "=ROUND(D30*(1+C35),1)", "0.0", "中性估值 ×（1+议价空间）", key=True)
out(37, "目标成交价(万元)", "=D30", "0.0", "中性估值", key=True)
out(38, "心理底价(万元)", "=D29", "0.0", "保守估值；低于此价需慎重", key=True)

# comps list (default criteria)
wv.cell(row=40, column=2, value="六、默认条件下的可比样本（修改条件后请到「成交明细」按 Y 列筛选）").font = f_bold
ch_ = ["成交日期", "面积(㎡)", "楼层段", "朝向", "首图观感", "总价(万元)", "单价(元/㎡)"]
src_cols = ["D", "N", "P", "R", "T", "U", "W"]
for i, h in enumerate(ch_, 2):
    c = wv.cell(row=41, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
comp_rows = []
for idx, rw in enumerate(rows, 1):
    rm = int(rw[3][0]); area = rw[4]; d = dt.datetime.strptime(rw[2], "%Y-%m-%d")
    if rm == 3 and 105 <= area <= 117 and d >= dt.datetime(2026, 3, 1):
        comp_rows.append(idx + 1)
for j, sr in enumerate(comp_rows):
    r = 42 + j
    for i, sc in enumerate(src_cols, 2):
        c = wv.cell(row=r, column=i, value=f"='{D}'!{sc}{sr}")
        c.font = f_link; c.border = border; c.alignment = center
    wv.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    wv.cell(row=r, column=3).number_format = "0.00"
    wv.cell(row=r, column=7).number_format = "0.0"
    wv.cell(row=r, column=8).number_format = "#,##0"
for i, w in enumerate([3, 26, 16, 16, 12, 12, 12, 12], 1):
    wv.column_dimensions[get_column_letter(i)].width = w
wv.column_dimensions["D"].width = 16
# long notes in column D overflow: widen E.. not needed, wrap off for notes
for r in list(range(5, 11)) + [22, 23, 24, 35, 36, 37, 38]:
    wv.cell(row=r, column=4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

# ---------------------------------------------------------------- 持有vs出售测算
wh = wb.create_sheet("持有vs出售测算")
wh["A1"] = "持有 vs 出售：N 年后的资产对比（简化模型）"
wh["A1"].font = f_title
wh["A2"] = "出售方案：现在按估值卖出，扣折扣与交易成本后，资金按年化收益率滚存。持有方案：N 年后按变动后的房价卖出，另加 N 年净租金（未计租金再投资）。若为自住，把「月租金」理解为自住省下的房租。"
wh["A2"].font = f_sub
wh["A2"].alignment = left
wh.merge_cells("A2:H2")
wh.row_dimensions[2].height = 32

def hin(r, label, val, fmt, note, link=False):
    wh.cell(row=r, column=1, value=label).font = f_base
    c = wh.cell(row=r, column=2, value=val)
    c.border = border; c.alignment = center; c.number_format = fmt
    if link:
        c.font = f_link
    else:
        c.font = f_input; c.fill = fill_input
    n = wh.cell(row=r, column=3, value=note); n.font = f_sub

wh.cell(row=4, column=1, value="参数").font = f_bold
hin(5, "当前估值(万元)", "='本案估值'!D30", "0.0", "引自「本案估值」中性估值", link=True)
hin(6, "本案面积(㎡)", "='本案估值'!C5", "0.00", "引自「本案估值」", link=True)
hin(7, "成交折扣（相对估值让价）", 0.02, "0.0%", "假设值：出售时通常需在目标价上再让一点")
hin(8, "卖方交易成本", 0.01, "0.0%", "假设值：中介/杂费等；满五唯一可免个税、满两年免增值税，请以税务口径核实")
hin(9, "对比年限 N（年）", 3, "0", "可改为 1~10")
hin(10, "月租金(元)", 3000, "#,##0", "假设值：请以贝壳/安居客同户型在租价格核实")
hin(11, "年空置(月)", 1, "0.0", "假设值")
hin(12, "物业费(元/㎡/月)", 2.8, "0.00", "来源：房天下小区页物业费 2.8 元/㎡·月")
hin(13, "年维修/家电折旧(元)", 3000, "#,##0", "假设值")
hin(14, "卖房资金年化收益率", 0.02, "0.0%", "假设值：低风险理财/国债/提前还房贷（有房贷时可用房贷利率）")

wh.cell(row=16, column=1, value="中间结果").font = f_bold
def hout(r, label, formula, fmt, note="", key=False):
    wh.cell(row=r, column=1, value=label).font = f_bold if key else f_base
    c = wh.cell(row=r, column=2, value=formula)
    c.font = f_bold if key else f_base; c.border = border; c.alignment = center; c.number_format = fmt
    if key: c.fill = fill_key
    wh.cell(row=r, column=3, value=note).font = f_sub
hout(17, "年净租金(万元)", "=(B10*(12-B11)-B12*B6*12-B13)/10000", "0.00", "毛租金 − 空置 − 物业费 − 维修")
hout(18, "净租金收益率", "=B17/B5", "0.00%", "相对当前估值")
hout(19, "出售方案 N 年后资金(万元)", "=B5*(1-B7)*(1-B8)*(1+B14)^B9", "0.0")
hout(20, "盈亏平衡：房价年变动率", "=((1+B14)^B9-B17*B9/(B5*(1-B7)*(1-B8)))^(1/B9)-1", "0.00%",
     "房价年均变动高于此值 → 持有更划算；低于此值 → 现在卖更划算", key=True)

wh.cell(row=22, column=1, value="情景对比").font = f_bold
sh = ["房价年变动率", "N 年后房价(万元)", "持有方案 N 年后资产(万元)", "出售方案 N 年后资金(万元)", "持有 − 出售(万元)", "结论"]
for i, h in enumerate(sh, 1):
    c = wh.cell(row=23, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
gs = [-0.06, -0.04, -0.02, 0.0, 0.02, 0.04]
for j, g in enumerate(gs):
    r = 24 + j
    c = wh.cell(row=r, column=1, value=g); c.font = f_input; c.fill = fill_input; c.number_format = "0.0%;-0.0%;0.0%"
    wh.cell(row=r, column=2, value=f"=$B$5*(1+A{r})^$B$9").number_format = "0.0"
    wh.cell(row=r, column=3, value=f"=B{r}*(1-$B$7)*(1-$B$8)+$B$17*$B$9").number_format = "0.0"
    wh.cell(row=r, column=4, value="=$B$19").number_format = "0.0"
    wh.cell(row=r, column=5, value=f"=C{r}-D{r}").number_format = "0.0;-0.0;0.0"
    wh.cell(row=r, column=6, value=f"=IF(E{r}>0,\"持有占优\",\"出售占优\")")
    for cc in range(1, 7):
        c = wh.cell(row=r, column=cc); c.border = border; c.alignment = center
        if cc > 1: c.font = f_base
wh["A31"] = "说明：模型未计入房贷利息（如有房贷，把 B14 改成房贷利率更贴近真实机会成本）、未计入换房/搬家成本和自住效用；只用于量化「房价要涨/跌多少才值得继续拿着」。"
wh["A31"].font = f_sub
wh.merge_cells("A31:F31")
wh["A31"].alignment = left
wh.row_dimensions[31].height = 30
for i, w in enumerate([30, 18, 26, 24, 18, 12], 1):
    wh.column_dimensions[get_column_letter(i)].width = w
wh.row_dimensions[23].height = 30

for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_view.showGridLines = False if sheet.title in ("说明", "本案估值", "持有vs出售测算") else True

wb.save("../data/国创光谷上城_成交数据与卖房参考.xlsx")
print("saved", len(comp_rows), "comps")
